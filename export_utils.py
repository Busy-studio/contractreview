from __future__ import annotations

import html
import io
import re
from datetime import datetime
from typing import Iterable, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


_HTML_TAG_RE = re.compile(r"<[^>]+>")
_MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)\]\([^\)]+\)")


def _strip_inline_markup(text: str) -> str:
    text = str(text or "")
    text = re.sub(r"<chg>(.*?)</chg>", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"<span[^>]*>(.*?)</span>", r"\1", text, flags=re.DOTALL | re.IGNORECASE)
    text = _HTML_TAG_RE.sub("", text)
    text = _MARKDOWN_LINK_RE.sub(r"\1", text)
    text = text.replace("**", "").replace("__", "").replace(chr(96), "")
    for icon in ("📌", "📄", "⚠️", "⚠", "⚖️", "⚖", "✏️", "✏", "🔎", "ℹ️", "ℹ"):
        text = text.replace(icon, "")
    text = html.unescape(text)
    return re.sub(r"[ \t]+", " ", text).strip()


def _iter_report_lines(report_text: str) -> Iterable[Tuple[str, str]]:
    for raw in str(report_text or "").splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            yield "spacer", ""
            continue
        if stripped in {"---", "***", "___"}:
            yield "spacer", ""
            continue

        if stripped.startswith("#### "):
            yield "h3", _strip_inline_markup(stripped[5:])
        elif stripped.startswith("### "):
            yield "h2", _strip_inline_markup(stripped[4:])
        elif stripped.startswith("## "):
            yield "h1", _strip_inline_markup(stripped[3:])
        elif stripped.startswith("# "):
            yield "h1", _strip_inline_markup(stripped[2:])
        elif stripped.startswith("- "):
            yield "bullet", _strip_inline_markup(stripped[2:])
        else:
            yield "body", _strip_inline_markup(stripped)


def _set_docx_font(style, name: str, size_pt: float, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def build_docx_bytes(report_text: str, title: str = "계약 검토 결과") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    styles = doc.styles
    _set_docx_font(styles["Normal"], "맑은 고딕", 10.5, False)
    _set_docx_font(styles["Title"], "맑은 고딕", 18, True)
    _set_docx_font(styles["Heading 1"], "맑은 고딕", 14, True)
    _set_docx_font(styles["Heading 2"], "맑은 고딕", 12, True)
    _set_docx_font(styles["Heading 3"], "맑은 고딕", 10.5, True)

    p = doc.add_paragraph(style="Title")
    p.add_run(title)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    meta_run.font.size = Pt(8.5)

    last_was_spacer = False
    for kind, text in _iter_report_lines(report_text):
        if kind == "spacer":
            if not last_was_spacer:
                doc.add_paragraph("")
            last_was_spacer = True
            continue

        last_was_spacer = False
        if not text:
            continue

        if kind == "h1":
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(text)
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(text)
        elif kind == "h3":
            p = doc.add_paragraph(style="Heading 3")
            p.add_run(text)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            _set_docx_font(p.style, "맑은 고딕", 10.5, False)
            p.add_run(text)
        else:
            p = doc.add_paragraph()
            p.add_run(text)

        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

    doc.add_paragraph("")
    footer = doc.add_paragraph()
    r = footer.add_run("※ 본 결과는 내부 검토 참고용이며 최종 법률자문을 대체하지 않습니다.")
    r.italic = True
    r.font.size = Pt(8.5)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _register_pdf_fonts() -> tuple[str, str]:
    # ReportLab 기본 Korean Unicode CID font.
    # 별도 폰트 파일 없이 Linux/Streamlit 환경에서도 한글 텍스트 레이어를 유지한다.
    body_font = "HYSMyeongJo-Medium"
    try:
        pdfmetrics.getFont(body_font)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(body_font))
    return body_font, body_font


def _pdf_safe(text: str) -> str:
    return html.escape(_strip_inline_markup(text), quote=False).replace("\n", "<br/>")


def build_pdf_bytes(report_text: str, title: str = "계약 검토 결과") -> bytes:
    body_font, heading_font = _register_pdf_fonts()
    out = io.BytesIO()

    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=title,
        author="PNU 계약/협약서 검토",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "KTitle",
        parent=styles["Title"],
        fontName=heading_font,
        fontSize=17,
        leading=22,
        alignment=TA_LEFT,
        spaceAfter=7 * mm,
    )
    h1 = ParagraphStyle(
        "KH1",
        parent=styles["Heading1"],
        fontName=heading_font,
        fontSize=13,
        leading=18,
        spaceBefore=5 * mm,
        spaceAfter=2.5 * mm,
    )
    h2 = ParagraphStyle(
        "KH2",
        parent=styles["Heading2"],
        fontName=heading_font,
        fontSize=11.5,
        leading=16,
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
    )
    h3 = ParagraphStyle(
        "KH3",
        parent=styles["Heading3"],
        fontName=heading_font,
        fontSize=10.5,
        leading=15,
        spaceBefore=3 * mm,
        spaceAfter=1.5 * mm,
    )
    body = ParagraphStyle(
        "KBody",
        parent=styles["BodyText"],
        fontName=body_font,
        fontSize=9.5,
        leading=15,
        spaceAfter=2 * mm,
    )
    bullet = ParagraphStyle(
        "KBullet",
        parent=body,
        leftIndent=4 * mm,
        firstLineIndent=-3 * mm,
        spaceAfter=1.5 * mm,
    )
    meta = ParagraphStyle(
        "KMeta",
        parent=body,
        fontSize=8,
        leading=11,
        spaceAfter=4 * mm,
    )

    story = [
        Paragraph(_pdf_safe(title), title_style),
        Paragraph(_pdf_safe(f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}"), meta),
    ]

    last_was_spacer = False
    for kind, text in _iter_report_lines(report_text):
        if kind == "spacer":
            if not last_was_spacer:
                story.append(Spacer(1, 2.5 * mm))
            last_was_spacer = True
            continue

        last_was_spacer = False
        if not text:
            continue

        safe = _pdf_safe(text)
        if kind == "h1":
            story.append(Paragraph(safe, h1))
        elif kind == "h2":
            story.append(Paragraph(safe, h2))
        elif kind == "h3":
            story.append(Paragraph(safe, h3))
        elif kind == "bullet":
            story.append(Paragraph("- " + safe, bullet))
        else:
            story.append(Paragraph(safe, body))

    story.append(Spacer(1, 4 * mm))
    story.append(
        Paragraph(
            _pdf_safe("※ 본 결과는 내부 검토 참고용이며 최종 법률자문을 대체하지 않습니다."),
            meta,
        )
    )

    doc.build(story)
    return out.getvalue()
